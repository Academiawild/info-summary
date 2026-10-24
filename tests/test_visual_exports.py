from __future__ import annotations

import shutil
import struct
import tempfile
import unittest
import zlib
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document

from info_summary.exporters import ResultExporter
from info_summary.models import (
    ProcessingArtifact,
    ProcessingRequest,
    ProcessingStep,
    SourceType,
    TranscriptSegment,
    VisualNote,
)
from info_summary.result_index import ResultIndex


def write_test_png(path: Path, width: int = 160, height: int = 90) -> None:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    rows = (b"\x00" + b"\x2d\x7a\xa5" * width) * height
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(rows))
        + chunk(b"IEND", b"")
    )


class VisualExportTests(unittest.TestCase):
    def request(self, formats=("docx", "md")) -> ProcessingRequest:
        return ProcessingRequest(
            "https://example.test/video", SourceType.VIDEO_URL,
            steps=(ProcessingStep.ORIGINAL, ProcessingStep.SUMMARY),
            export_formats=formats, include_visual=True,
        )

    def artifact(self, image: Path) -> ProcessingArtifact:
        return ProcessingArtifact(
            title="演示视频", source="https://example.test/video", method="embedded_subtitle", language="zh",
            segments=[TranscriptSegment(0, 1, "这是讲述原文。")], original_text="这是讲述原文。",
            summary_text="[00:05] 图表显示增长，依据画面笔记。",
            visual_notes=[VisualNote(5, image, "蓝色图表显示增长。", end=8)],
            visual_overview="主要画面包含一张图表。",
        )

    def test_both_formats_embed_images_and_register_movable_relative_paths(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "原图 with space.png"
            write_test_png(source)
            result = ResultExporter().export(self.artifact(source), self.request(), root, "演示")
            self.assertEqual(
                set(result.output_files),
                {"original.docx", "original.md", "summary.docx", "summary.md", "visual.docx", "visual.md", "image.frame_0001.png"},
            )
            self.assertEqual(result.warnings, [])
            docx_path = result.output_files["visual.docx"]
            document = Document(docx_path)
            self.assertEqual(len(document.inline_shapes), 1)
            shape = document.inline_shapes[0]
            section = document.sections[0]
            self.assertLessEqual(shape.width, section.page_width - section.left_margin - section.right_margin)
            self.assertAlmostEqual(shape.width / shape.height, 160 / 90, places=5)
            with ZipFile(docx_path) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
                self.assertEqual(len(media), 1)
                self.assertEqual(archive.read(media[0]), source.read_bytes())
            markdown = result.output_files["visual.md"].read_text(encoding="utf-8")
            self.assertIn("[00:05 - 00:08 的画面](images/frame_0001.png)", markdown)
            self.assertIn("蓝色图表显示增长。", markdown)
            self.assertNotIn("蓝色图表", result.output_files["original.md"].read_text(encoding="utf-8"))
            moved = root / "moved"
            shutil.copytree(result.result_dir, moved)
            source.unlink()
            self.assertTrue((moved / "images/frame_0001.png").is_file())
            self.assertFalse(any(result.result_dir.rglob("*.tmp*")))

    def test_portrait_image_keeps_ratio_and_fits_the_page(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "portrait.png"
            write_test_png(source, 90, 240)
            result = ResultExporter().export(self.artifact(source), self.request(("docx",)), root, "演示")
            document = Document(result.output_files["visual.docx"])
            shape = document.inline_shapes[0]
            section = document.sections[0]
            self.assertLess(shape.height, section.page_height - section.top_margin - section.bottom_margin)
            self.assertAlmostEqual(shape.width / shape.height, 90 / 240, places=5)

    def test_inline_references_become_readable_times_in_word_and_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "frame.png"
            write_test_png(image)
            artifact = self.artifact(image)
            artifact.summary_text = "观点。\n\n[[frame:1]]\n\n解释。[[frame:99]]"
            result = ResultExporter().export(artifact, self.request(), root, "演示")
            markdown = result.output_files["summary.md"].read_text()
            word = "\n".join(p.text for p in Document(result.output_files["summary.docx"]).paragraphs)
            for text in (markdown, word):
                self.assertIn("（画面 00:05）", text)
                self.assertIn("（对应画面未取得）", text)
                self.assertNotIn("[[frame:", text)

    def test_repeated_source_image_is_copied_only_once(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.png"
            write_test_png(source)
            artifact = self.artifact(source)
            artifact.visual_notes.append(VisualNote(12, source, "同一页面的另一处说明。"))
            result = ResultExporter().export(artifact, self.request(("md",)), root, "演示")
            self.assertEqual(len([key for key in result.output_files if key.startswith("image.")]), 1)
            self.assertEqual(result.output_files["visual.md"].read_text(encoding="utf-8").count("](images/frame_0001.png)"), 2)

    def test_image_copy_failure_preserves_all_text_outputs_and_notes(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.png"
            write_test_png(source)
            with patch("info_summary.exporters.shutil.copyfile", side_effect=OSError("disk unavailable")):
                result = ResultExporter().export(self.artifact(source), self.request(), root, "演示")
            self.assertIn("original.docx", result.output_files)
            self.assertIn("summary.md", result.output_files)
            self.assertTrue(result.warnings)
            self.assertNotIn("image.frame_0001.png", result.output_files)
            self.assertIn("蓝色图表显示增长。", result.output_files["visual.md"].read_text(encoding="utf-8"))
            self.assertEqual(len(Document(result.output_files["visual.docx"]).inline_shapes), 0)

    def test_missing_or_invalid_image_keeps_description_with_warning(self) -> None:
        for invalid in (False, True):
            with self.subTest(invalid=invalid), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root / "source.png"
                if invalid:
                    source.write_bytes(b"not an image")
                result = ResultExporter().export(self.artifact(source), self.request(), root, "演示")
                self.assertTrue(result.warnings)
                self.assertIn("蓝色图表显示增长。", result.output_files["visual.md"].read_text(encoding="utf-8"))

    def test_missing_original_has_explicit_message_in_both_formats(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.png"
            write_test_png(source)
            artifact = self.artifact(source)
            artifact.original_text = ""
            artifact.segments = []
            result = ResultExporter().export(artifact, self.request(), root, "演示")
            original_docx = "\n".join(p.text for p in Document(result.output_files["original.docx"]).paragraphs)
            original_md = result.output_files["original.md"].read_text(encoding="utf-8")
            for text in (original_docx, original_md):
                self.assertIn("没有可用的字幕或语音原文", text)
                self.assertNotIn("蓝色图表", text)

    def test_empty_visual_result_still_writes_explanation(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            artifact = self.artifact(root / "missing.png")
            artifact.visual_notes = []
            result = ResultExporter().export(artifact, self.request(), root, "演示")
            self.assertTrue(result.warnings)
            self.assertIn("未取得可用的画面截图", result.output_files["visual.md"].read_text(encoding="utf-8"))
            self.assertTrue(result.output_files["visual.docx"].is_file())

    def test_disabled_visual_exports_keep_existing_behavior(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            request = replace(self.request(), include_visual=False)
            result = ResultExporter().export(self.artifact(root / "missing.png"), request, root, "演示")
            self.assertEqual(set(result.output_files), {"original.docx", "original.md", "summary.docx", "summary.md"})
            self.assertEqual(result.warnings, [])

    def test_reuse_requires_visual_documents_and_every_image(self) -> None:
        for missing in (None, "visual.docx", "visual.md", "image.frame_0001.png"):
            with self.subTest(missing=missing), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root / "source.png"
                write_test_png(source)
                request = self.request()
                result = ResultExporter().export(self.artifact(source), request, root, "演示")
                index = ResultIndex(root / "index.json")
                index.register(request, root, result)
                if missing:
                    result.output_files[missing].unlink()
                    self.assertIsNone(index.find(request, root))
                else:
                    self.assertEqual(index.find(request, root).output_files, result.output_files)

    def test_visual_index_without_any_registered_image_is_not_reused(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.png"
            write_test_png(source)
            request = self.request()
            result = ResultExporter().export(self.artifact(source), request, root, "演示")
            result.output_files.pop("image.frame_0001.png")
            index = ResultIndex(root / "index.json")
            index.register(request, root, result)
            self.assertIsNone(index.find(request, root))


if __name__ == "__main__":
    unittest.main()
