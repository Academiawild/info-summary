from __future__ import annotations

import tempfile
import threading
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document

from info_summary.audio.audio_transcriber import AudioTranscriber
from info_summary.exporters import ResultExporter, format_timestamp
from info_summary.video.subtitles import parse_subtitle_file
from info_summary.macos import app_data_dir, require_macos
from info_summary.models import (
    DownloadOptions,
    OutputLanguage,
    ProcessingArtifact,
    ProcessingRequest,
    ProcessingResult,
    ProcessingStep,
    SourceType,
    TranscriptSegment,
    VideoDownloadRequest,
    VideoDownloadResult,
    VideoQuality,
)
from info_summary.pipeline import ProcessingPipeline
from info_summary.result_index import ResultIndex
from info_summary.service import ProcessingService
from info_summary.sources import detect_source, is_supported_video_url
from info_summary.text_pipeline import ResumableTextProcessor
from info_summary.ui import TaskView
from info_summary.video.video_extractor import VideoExtractor
from info_summary.video.video_downloader import VideoDownloadIndex, VideoDownloader, VideoDownloadService


class FakeProvider:
    def __init__(self, healthy: bool = True) -> None:
        self.healthy = healthy
        self.calls: list[tuple[str, str]] = []

    def health_check(self):
        return (self.healthy, "模型不可用")

    def list_models(self):
        return ["fake"]

    def generate(self, model: str, system_prompt: str, user_prompt: str) -> str:
        self.calls.append((system_prompt, user_prompt))
        if "总结" in system_prompt or "摘要" in system_prompt:
            return "总结内容"
        current = user_prompt.split("<current>只整理下面的当前内容：</current>\n")[-1]
        return "润色内容：" + current


class ExportTests(unittest.TestCase):
    def artifact(self) -> ProcessingArtifact:
        return ProcessingArtifact(
            title="测试视频", source="https://example.test/video", method="embedded_subtitle",
            language="zh", segments=[TranscriptSegment(0, 2, "第一句。"), TranscriptSegment(3, 5, "第二句。")],
            original_text="第一句。\n\n第二句。", polished_text="润色后的内容。", summary_text="总结内容。",
        )

    def test_every_selected_step_and_format_gets_its_own_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for formats in (("docx",), ("md",), ("docx", "md")):
                request = ProcessingRequest(
                    "https://example.test/video", SourceType.VIDEO_URL,
                    steps=(ProcessingStep.ORIGINAL, ProcessingStep.POLISH, ProcessingStep.SUMMARY),
                    export_formats=formats,
                )
                result = ResultExporter().export(self.artifact(), request, root, "测试视频")
                self.assertEqual(len(result.output_files), 3 * len(formats))
                self.assertTrue(all(path.exists() for path in result.output_files.values()))

    def test_original_word_is_clean_text_and_markdown_has_timestamps(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            request = ProcessingRequest(
                "source", SourceType.VIDEO_URL,
                steps=(ProcessingStep.ORIGINAL,), export_formats=("docx", "md"),
            )
            result = ResultExporter().export(self.artifact(), request, Path(temp), "标题")
            paragraphs = [p.text for p in Document(result.output_files["original.docx"]).paragraphs]
            markdown = result.output_files["original.md"].read_text(encoding="utf-8")
            self.assertIn("第一句。", paragraphs)
            self.assertNotIn("[00:00", "\n".join(paragraphs))
            self.assertIn("[00:00 - 00:02]", markdown)

    def test_word_summary_uses_readable_bullet_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            artifact = self.artifact()
            artifact.summary_text = "全文主旨。\n\n- 第一点。\n- 第二点。"
            artifact.original_text = "- 原文中的符号保留。"
            request = ProcessingRequest(
                "source", SourceType.VIDEO_URL,
                steps=(ProcessingStep.ORIGINAL, ProcessingStep.SUMMARY), export_formats=("docx",),
            )
            result = ResultExporter().export(artifact, request, Path(temp), "标题")
            summary = Document(result.output_files["summary.docx"])
            bullets = [paragraph.text for paragraph in summary.paragraphs if paragraph.style.name == "List Bullet"]
            self.assertEqual(bullets, ["第一点。", "第二点。"])
            original = Document(result.output_files["original.docx"])
            self.assertIn(artifact.original_text, [paragraph.text for paragraph in original.paragraphs])

    def test_result_folder_names_do_not_collide(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            request = ProcessingRequest("source", SourceType.VIDEO_URL, steps=(ProcessingStep.ORIGINAL,), export_formats=("md",))
            first = ResultExporter().export(self.artifact(), request, Path(temp), "标题")
            second = ResultExporter().export(self.artifact(), request, Path(temp), "标题")
            self.assertNotEqual(first.result_dir, second.result_dir)

    def test_timestamp(self) -> None:
        self.assertEqual(format_timestamp(3661), "01:01:01")


class PipelineTests(unittest.TestCase):
    def make_pipeline(self, root: Path):
        transcriber = AudioTranscriber(root)
        extractor = VideoExtractor()
        return ProcessingPipeline(root, checkpoint_dir=root / "checkpoints", audio_transcriber=transcriber, video_extractor=extractor), transcriber, extractor

    def test_subtitle_only_does_not_touch_audio_whisper_or_provider(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pipeline, transcriber, extractor = self.make_pipeline(root)
            extractor.extract_info_with_cookie_fallback = MagicMock(return_value=({"title": "Video"}, DownloadOptions()))
            extractor.try_extract_subtitles = MagicMock(return_value=([TranscriptSegment(0, 1, "字幕")], "zh", "embedded_subtitle"))
            transcriber.check_audio_dependencies = MagicMock()
            transcriber.transcribe_audio = MagicMock()
            pipeline._provider_for = MagicMock(side_effect=AssertionError("provider created"))
            result = pipeline.run(ProcessingRequest(
                "https://vimeo.com/1", SourceType.VIDEO_URL, output_dir=str(root / "out"),
                steps=(ProcessingStep.ORIGINAL,), export_formats=("md",),
            ))
            self.assertTrue(result.output_files["original.md"].exists())
            transcriber.check_audio_dependencies.assert_not_called()
            transcriber.transcribe_audio.assert_not_called()
            pipeline._provider_for.assert_not_called()

    def test_video_without_subtitles_downloads_and_transcribes_audio(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pipeline, transcriber, extractor = self.make_pipeline(root)
            extractor.extract_info_with_cookie_fallback = MagicMock(return_value=({"title": "Video", "formats": [{"acodec": "aac"}]}, DownloadOptions()))
            extractor.try_extract_subtitles = MagicMock(return_value=([], "", ""))
            downloaded = root / "download.m4a"
            downloaded.touch()
            extractor.download_audio = MagicMock(return_value=downloaded)
            transcriber.check_audio_dependencies = MagicMock()
            transcriber.prepare_audio = MagicMock(return_value=downloaded)
            transcriber.transcribe_audio = MagicMock(return_value=[TranscriptSegment(0, 1, "转写")])
            transcriber.last_detected_language = "zh"
            pipeline.run(ProcessingRequest(
                "https://vimeo.com/1", SourceType.VIDEO_URL, output_dir=str(root / "out"),
                steps=(ProcessingStep.ORIGINAL,), export_formats=("docx",),
            ))
            extractor.download_audio.assert_called_once()
            transcriber.transcribe_audio.assert_called_once()

    def test_summary_without_polish_uses_original(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.touch()
            pipeline, transcriber, _ = self.make_pipeline(root)
            transcriber.check_audio_dependencies = MagicMock()
            transcriber.prepare_audio = MagicMock(return_value=audio)
            transcriber.transcribe_audio = MagicMock(return_value=[TranscriptSegment(0, 1, "原始内容")])
            provider = FakeProvider()
            pipeline._provider_for = MagicMock(return_value=provider)
            result = pipeline.run(ProcessingRequest(
                str(audio), SourceType.AUDIO_FILE, output_dir=str(root / "out"),
                steps=(ProcessingStep.ORIGINAL, ProcessingStep.SUMMARY), export_formats=("md",),
                text_model_name="fake",
            ))
            self.assertIn("原始内容", provider.calls[0][1])
            self.assertIn("summary.md", result.output_files)
            self.assertNotIn("polish.md", result.output_files)

    def test_polish_and_summary_use_one_provider(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.touch()
            pipeline, transcriber, _ = self.make_pipeline(root)
            transcriber.check_audio_dependencies = MagicMock()
            transcriber.prepare_audio = MagicMock(return_value=audio)
            transcriber.transcribe_audio = MagicMock(return_value=[TranscriptSegment(0, 1, "原始内容")])
            provider = FakeProvider()
            pipeline._provider_for = MagicMock(return_value=provider)
            result = pipeline.run(ProcessingRequest(
                str(audio), SourceType.AUDIO_FILE, output_dir=str(root / "out"),
                export_formats=("md",), text_model_name="fake",
            ))
            self.assertIn("polish.md", result.output_files)
            self.assertIn("summary.md", result.output_files)
            pipeline._provider_for.assert_called_once()

    def test_manually_entered_model_can_generate_without_model_list_access(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pipeline, _, _ = self.make_pipeline(root)
            pipeline._obtain_original = MagicMock(return_value=(
                "title", [TranscriptSegment(0, 1, "原始内容")], "subtitle", "zh",
            ))
            provider = FakeProvider(healthy=False)
            provider.health_check = MagicMock(side_effect=AssertionError("model list is unavailable"))
            provider.list_models = MagicMock(side_effect=AssertionError("model list is unavailable"))
            pipeline._provider_for = MagicMock(return_value=provider)
            result = pipeline.run(ProcessingRequest(
                "https://vimeo.com/1", SourceType.VIDEO_URL, output_dir=str(root / "out"),
                export_formats=("md",), text_model_name="manually-entered-model",
            ))
            self.assertTrue(result.output_files["polish.md"].is_file())
            self.assertTrue(result.output_files["summary.md"].is_file())
            self.assertEqual(result.warnings, [])
            self.assertTrue(provider.calls)
            provider.health_check.assert_not_called()
            provider.list_models.assert_not_called()

    def test_missing_subtitle_and_audio_reports_clear_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pipeline, _, extractor = self.make_pipeline(root)
            extractor.extract_info_with_cookie_fallback = MagicMock(return_value=({"title": "Empty", "formats": []}, DownloadOptions()))
            extractor.try_extract_subtitles = MagicMock(return_value=([], "", ""))
            with self.assertRaisesRegex(RuntimeError, "没有可用字幕"):
                pipeline.run(ProcessingRequest("https://vimeo.com/1", SourceType.VIDEO_URL, output_dir=str(root / "out")))

    def test_cancelled_task_exports_nothing(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.touch()
            event = threading.Event()
            event.set()
            with self.assertRaisesRegex(RuntimeError, "已取消"):
                ProcessingPipeline(root, checkpoint_dir=root / "checkpoints").run(ProcessingRequest(str(audio), SourceType.AUDIO_FILE), cancel_event=event)


class IndexAndServiceTests(unittest.TestCase):
    def test_index_finds_only_existing_completed_result(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.write_bytes(b"audio")
            output = root / "result"
            output.mkdir()
            file = output / "原文稿.md"
            file.write_text("ok", encoding="utf-8")
            request = ProcessingRequest(str(audio), SourceType.AUDIO_FILE, steps=(ProcessingStep.ORIGINAL,), export_formats=("md",))
            index = ResultIndex(root / "results.json")
            result = ProcessingResult("title", str(audio), "audio", "zh", output, {"original.md": file})
            index.register(request, root, result)
            self.assertEqual(index.find(request, root).result_dir, output)
            file.unlink()
            self.assertIsNone(index.find(request, root))

    def test_settings_or_steps_change_duplicate_key(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.write_bytes(b"audio")
            original = ProcessingRequest(str(audio), SourceType.AUDIO_FILE, steps=(ProcessingStep.ORIGINAL,))
            summary = ProcessingRequest(str(audio), SourceType.AUDIO_FILE, steps=(ProcessingStep.ORIGINAL, ProcessingStep.SUMMARY))
            self.assertNotEqual(original.result_key(root), summary.result_key(root))

    def test_audio_folder_expands_to_one_request_per_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "audio"
            folder.mkdir()
            (folder / "a.mp3").touch()
            (folder / "b.wav").touch()
            (folder / "ignore.txt").touch()
            service = ProcessingService(root, result_index=ResultIndex(root / "results.json"), pipeline=MagicMock())
            requests = service.expand_requests(ProcessingRequest(str(folder), SourceType.AUDIO_FOLDER))
            self.assertEqual(len(requests), 2)
            self.assertTrue(all(item.source_type == SourceType.AUDIO_FILE for item in requests))

    def test_service_registers_only_after_success(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.touch()
            index = MagicMock()
            pipeline = MagicMock()
            pipeline.run.side_effect = RuntimeError("failed")
            service = ProcessingService(root, result_index=index, pipeline=pipeline)
            with self.assertRaises(RuntimeError):
                service.run(ProcessingRequest(str(audio), SourceType.AUDIO_FILE))
            index.register.assert_not_called()

    def test_forced_rerun_replaces_the_latest_index_result(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            audio = root / "audio.wav"
            audio.touch()
            request = ProcessingRequest(
                str(audio), SourceType.AUDIO_FILE,
                steps=(ProcessingStep.ORIGINAL,), export_formats=("md",),
            )
            index = ResultIndex(root / "results.json")
            old_dir = root / "old"
            new_dir = root / "new"
            old_dir.mkdir()
            new_dir.mkdir()
            old_file = old_dir / "原文稿.md"
            new_file = new_dir / "原文稿.md"
            old_file.touch()
            new_file.touch()
            index.register(request, root, ProcessingResult("old", str(audio), "audio", "zh", old_dir, {"original.md": old_file}))
            pipeline = MagicMock()
            pipeline.run.return_value = ProcessingResult("new", str(audio), "audio", "zh", new_dir, {"original.md": new_file})
            ProcessingService(root, result_index=index, pipeline=pipeline).run(request)
            self.assertEqual(index.find(request, root).result_dir, new_dir)


class SourceAndMacTests(unittest.TestCase):
    def test_supported_sources(self) -> None:
        self.assertTrue(is_supported_video_url("https://www.youtube.com/watch?v=x"))
        self.assertFalse(is_supported_video_url("https://youtube.com.evil.test/watch?v=x"))
        with tempfile.NamedTemporaryFile(suffix=".mp3") as audio:
            self.assertEqual(detect_source(audio.name, Path.cwd())[0], SourceType.AUDIO_FILE)

    def test_macos_data_directory(self) -> None:
        with patch.dict("os.environ", {"XDG_DATA_HOME": "/tmp/custom"}):
            self.assertEqual(app_data_dir(), Path("/tmp/custom/InfoSummary"))

    def test_non_macos_is_rejected(self) -> None:
        with patch("info_summary.macos.sys.platform", "win32"):
            with self.assertRaisesRegex(RuntimeError, "只支持 macOS"):
                require_macos()

    def test_subtitle_parser_and_language_conversion(self) -> None:
        content = """WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nhello <b>world</b>\n"""
        with tempfile.NamedTemporaryFile("w", suffix=".vtt", encoding="utf-8", delete=False) as handle:
            handle.write(content)
            path = Path(handle.name)
        try:
            segments = parse_subtitle_file(path)
        finally:
            path.unlink(missing_ok=True)
        self.assertEqual(segments[0].text, "hello world")
        self.assertEqual(ResumableTextProcessor().apply_output_language("內容總結", OutputLanguage.SIMPLIFIED_CHINESE), "内容总结")


class InterfaceTests(unittest.TestCase):
    def test_step_and_format_choices_do_not_use_x_checkboxes(self) -> None:
        try:
            root = tk.Tk()
        except tk.TclError as exc:
            self.skipTest(str(exc))
        root.withdraw()
        callbacks = {
            name: (lambda: None)
            for name in (
                "choose_audio", "choose_folder", "choose_output", "choose_cookies",
                "check_provider", "load_models", "start_ollama", "start", "cancel", "open_result",
            )
        }
        try:
            view = TaskView(root, callbacks)
            classes: list[str] = []

            def collect(widget) -> None:
                classes.append(widget.winfo_class())
                for child in widget.winfo_children():
                    collect(child)

            collect(root)
            self.assertNotIn("TCheckbutton", classes)
            view.preset_choice.set("整理成文稿")
            view.preset_combo.event_generate("<<ComboboxSelected>>")
            self.assertTrue(view.include_polish.get())
            self.assertFalse(view.include_summary.get())
            self.assertIn("润色稿", view.output_preview.get())
            view.select_mode("download")
            self.assertEqual(view.start_button.cget("text"), "开始下载")
            self.assertEqual(view.process_steps_card.winfo_manager(), "")
            self.assertEqual(view.formats.winfo_manager(), "")
        finally:
            root.destroy()


class VideoDownloadTests(unittest.TestCase):
    def test_downloads_video_and_audio_as_mp4_with_progress(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            request = VideoDownloadRequest(
                "https://vimeo.com/123",
                output_dir=str(root),
                quality=VideoQuality.HD_1080,
            )
            downloader = VideoDownloader()
            downloader.client.extract_info_with_cookie_fallback = MagicMock(
                return_value=({"title": "Demo Video", "id": "123"}, DownloadOptions())
            )
            captured: dict = {}

            class FakeYDL:
                def __init__(self, options):
                    captured.update(options)

                def __enter__(self):
                    return self

                def __exit__(self, *_args):
                    return False

                def download(self, _urls):
                    target = Path(captured["outtmpl"].replace("%(ext)s", "mp4"))
                    target.write_bytes(b"video")
                    captured["progress_hooks"][0](
                        {"status": "downloading", "total_bytes": 100, "downloaded_bytes": 50}
                    )
                    captured["progress_hooks"][0]({"status": "finished"})

            progress: list[float] = []
            with patch("info_summary.video.video_downloader.YoutubeDL", FakeYDL):
                result = downloader.download(request, progress=lambda value, _message: progress.append(value))
            self.assertTrue(result.video_path.exists())
            self.assertEqual(result.video_path.suffix, ".mp4")
            self.assertIn("height<=1080", captured["format"])
            self.assertEqual(captured["merge_output_format"], "mp4")
            self.assertEqual(captured["postprocessors"][0]["preferedformat"], "mp4")
            self.assertIn(0.5, [round(value, 2) for value in progress])
            self.assertEqual(progress[-1], 1.0)

    def test_download_index_and_forced_redownload(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            request = VideoDownloadRequest("https://vimeo.com/123", quality=VideoQuality.BEST)
            old_dir = root / "old"
            new_dir = root / "new"
            old_dir.mkdir()
            new_dir.mkdir()
            old_video = old_dir / "视频.mp4"
            new_video = new_dir / "视频.mp4"
            old_video.touch()
            new_video.touch()
            index = VideoDownloadIndex(root / "downloads.json")
            old = VideoDownloadResult("old", request.url, request.quality, old_dir, old_video)
            index.register(request, old)
            self.assertEqual(index.find(request).video_path, old_video)
            downloader = MagicMock()
            downloader.download.return_value = VideoDownloadResult(
                "new", request.url, request.quality, new_dir, new_video
            )
            VideoDownloadService(downloader=downloader, index=index).run(request)
            self.assertEqual(index.find(request).video_path, new_video)

    def test_cancelled_download_is_not_registered(self) -> None:
        request = VideoDownloadRequest("https://vimeo.com/123")
        downloader = MagicMock()
        downloader.download.side_effect = RuntimeError("视频下载已取消")
        index = MagicMock()
        service = VideoDownloadService(downloader=downloader, index=index)
        with self.assertRaisesRegex(RuntimeError, "已取消"):
            service.run(request)
        index.register.assert_not_called()


if __name__ == "__main__":
    unittest.main()
