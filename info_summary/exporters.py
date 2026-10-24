from __future__ import annotations

import os
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .models import ProcessingArtifact, ProcessingRequest, ProcessingResult, ProcessingStep, VisualNote
from .formatting import format_timestamp, readable_original, readable_summary
from .html_report import REPORT_FILENAME, render_report
from .output_paths import create_result_dir


STEP_NAMES = {
    ProcessingStep.ORIGINAL: "原文稿",
    ProcessingStep.POLISH: "润色稿",
    ProcessingStep.SUMMARY: "内容总结",
}
NO_ORIGINAL_TEXT = "没有可用的字幕或语音原文。"
VisualImage = tuple[VisualNote, Path | None]


class ResultExporter:
    def export(
        self,
        artifact: ProcessingArtifact,
        request: ProcessingRequest,
        output_root: Path,
        safe_title: str,
    ) -> ProcessingResult:
        result_dir = self._create_result_dir(output_root, safe_title)
        output_files: dict[str, Path] = {}
        for step, text in ((ProcessingStep.POLISH, artifact.polished_text),
                           (ProcessingStep.SUMMARY, artifact.summary_text)):
            if request.has_step(step) and not text.strip():
                warning = f"{STEP_NAMES[step]}尚未生成完整内容，请修复后继续未完成任务。"
                if warning not in artifact.warnings:
                    artifact.warnings.append(warning)
        visual_images: list[VisualImage] = []
        if request.include_visual:
            visual_images = self._copy_visual_images(artifact, result_dir, output_files)
        # Prepare the embedded images first so any reader warnings also appear
        # in the optional Word/Markdown files generated below.
        report = None
        if "html" in request.export_formats:
            try:
                report = render_report(artifact, request, visual_images)
            except (OSError, UnicodeError, ValueError) as exc:
                artifact.warnings.append(f"阅读报告未能生成，其他已保存文件仍可使用：{exc}")
        report_warnings = list(artifact.warnings)

        def save(key: str, path: Path, write) -> None:
            try:
                write()
            except (OSError, UnicodeError, ValueError) as exc:
                artifact.warnings.append(f"{path.name} 未能保存，已保留其他可用结果：{exc}")
            else:
                output_files[key] = path
        content = {
            ProcessingStep.ORIGINAL: artifact.original_text if artifact.original_text.strip() else NO_ORIGINAL_TEXT,
            ProcessingStep.POLISH: artifact.polished_text,
            ProcessingStep.SUMMARY: readable_summary(artifact.summary_text, artifact.visual_notes),
        }
        for step in request.steps:
            text = content.get(step, "")
            if step != ProcessingStep.ORIGINAL and not text.strip():
                continue
            name = STEP_NAMES[step]
            if "docx" in request.export_formats:
                path = result_dir / f"{name}.docx"
                save(f"{step.value}.docx", path, lambda: self._write_docx(path, artifact, name, text))
            if "md" in request.export_formats:
                path = result_dir / f"{name}.md"
                if step == ProcessingStep.ORIGINAL:
                    save(f"{step.value}.md", path, lambda: self._write_original_markdown(path, artifact))
                else:
                    save(f"{step.value}.md", path, lambda: self._write_text_markdown(path, artifact, name, text))
        if request.include_visual:
            if "docx" in request.export_formats:
                path = result_dir / "画面笔记.docx"
                save("visual.docx", path, lambda: self._write_visual_docx(path, artifact, visual_images))
            if "md" in request.export_formats:
                path = result_dir / "画面笔记.md"
                save("visual.md", path, lambda: self._write_visual_markdown(path, artifact, visual_images))
        if report is not None:
            path = result_dir / REPORT_FILENAME
            if artifact.warnings != report_warnings:
                try:
                    report = render_report(artifact, request, visual_images)
                except (OSError, UnicodeError, ValueError) as exc:
                    artifact.warnings.append(f"阅读报告处理提示未能更新：{exc}")
            save("report.html", path, lambda: self._write_text_atomic(path, report))
        if not output_files:
            raise RuntimeError("没有结果文件保存成功。" + "；".join(artifact.warnings))
        return ProcessingResult(
            title=artifact.title,
            source=artifact.source,
            method=artifact.method,
            language=artifact.language,
            result_dir=result_dir,
            output_files=output_files,
            warnings=artifact.warnings,
        )

    def _create_result_dir(self, output_root: Path, safe_title: str) -> Path:
        return create_result_dir(output_root, safe_title)

    def _write_docx(self, path: Path, artifact: ProcessingArtifact, heading: str, text: str) -> None:
        document = self._new_document(artifact, heading)
        self._add_text_blocks(document, text, summary=heading == STEP_NAMES[ProcessingStep.SUMMARY])
        self._save_docx_atomic(document, path)

    def _new_document(self, artifact: ProcessingArtifact, heading: str) -> Document:
        document = Document()
        self._set_default_font(document)
        document.add_heading(f"{artifact.title} - {heading}", level=1)
        self._paragraph(document, f"来源：{artifact.source}")
        self._paragraph(document, f"生成方式：{artifact.method}")
        self._paragraph(document, f"识别语言：{artifact.language}")
        self._paragraph(document, f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        if artifact.warnings:
            document.add_heading("处理提示", level=2)
            for warning in artifact.warnings:
                self._paragraph(document, warning)
        document.add_heading(heading, level=2)
        return document

    def _add_text_blocks(self, document: Document, text: str, *, summary: bool = False) -> None:
        for block in re.split(r"\n\s*\n", text):
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if summary and lines and all(line.startswith("- ") for line in lines):
                for line in lines:
                    self._paragraph(document, line[2:].strip(), style="List Bullet")
            elif block.strip():
                self._paragraph(document, block.strip())

    def _save_docx_atomic(self, document: Document, path: Path) -> None:
        temp = path.with_name(path.stem + ".tmp.docx")
        try:
            document.save(temp)
            os.replace(temp, path)
        finally:
            temp.unlink(missing_ok=True)

    def _write_original_markdown(self, path: Path, artifact: ProcessingArtifact) -> None:
        lines = self._markdown_header(artifact, "原文稿")
        for segment in artifact.segments:
            text = re.sub(r"\s+", " ", segment.text).strip()
            if text:
                lines.extend([f"**[{format_timestamp(segment.start)} - {format_timestamp(segment.end)}]** {text}", ""])
        if not any(segment.text.strip() for segment in artifact.segments):
            lines.extend([artifact.original_text.strip() or NO_ORIGINAL_TEXT, ""])
        self._write_text_atomic(path, "\n".join(lines).rstrip() + "\n")

    def _copy_visual_images(
        self, artifact: ProcessingArtifact, result_dir: Path, output_files: dict[str, Path],
    ) -> list[VisualImage]:
        """Keep each exported image with its note; a failed image must not discard text."""
        images: list[VisualImage] = []
        copied: dict[Path, Path | None] = {}
        for note in artifact.visual_notes:
            source = Path(note.image_path)
            if source not in copied:
                suffix = source.suffix.lower()
                if not re.fullmatch(r"\.[a-z0-9]{1,8}", suffix):
                    suffix = ".png"
                target = result_dir / "images" / f"frame_{len(copied) + 1:04d}{suffix}"
                temp = target.with_suffix(target.suffix + ".tmp")
                try:
                    # Validate once so Word and Markdown both get usable image files.
                    Image.from_file(str(source))
                    target.parent.mkdir(exist_ok=True)
                    shutil.copyfile(source, temp)
                    os.replace(temp, target)
                except Exception:
                    copied[source] = None
                    artifact.warnings.append(f"{format_timestamp(note.timestamp)} 的画面截图未能保存，已保留文字解读。")
                else:
                    copied[source] = target
                    output_files[f"image.{target.name}"] = target
                finally:
                    try:
                        temp.unlink(missing_ok=True)
                    except OSError:
                        pass
            images.append((note, copied[source]))
        if not artifact.visual_notes:
            artifact.warnings.append("画面分析未生成可用截图，画面笔记记录了当前处理情况。")
        return images

    def _visual_time_label(self, note: VisualNote) -> str:
        start = format_timestamp(note.timestamp)
        return f"{start} - {format_timestamp(note.end)}" if note.end is not None else start

    def _write_visual_docx(self, path: Path, artifact: ProcessingArtifact, images: list[VisualImage]) -> None:
        document = self._new_document(artifact, "画面笔记")
        self._add_text_blocks(document, artifact.visual_overview.strip() or "以下截图和解读用于补充理解视频中的画面。")
        section = document.sections[0]
        max_width = section.page_width - section.left_margin - section.right_margin
        max_height = min(Inches(5.5), section.page_height - section.top_margin - section.bottom_margin - Inches(1))
        for note, image_path in images:
            document.add_heading(self._visual_time_label(note), level=2)
            if image_path is not None:
                image = Image.from_file(str(image_path))
                scale = min(max_width / image.px_width, max_height / image.px_height)
                document.add_picture(str(image_path), width=int(image.px_width * scale), height=int(image.px_height * scale))
            else:
                self._paragraph(document, "截图未能保存，以下仅保留已经取得的文字解读。")
            self._add_text_blocks(document, note.description.strip() or "此画面尚未取得可靠解读。")
        if not images:
            self._paragraph(document, "未取得可用的画面截图。请查看处理提示，修复后可以重试。")
        self._save_docx_atomic(document, path)

    def _write_visual_markdown(self, path: Path, artifact: ProcessingArtifact, images: list[VisualImage]) -> None:
        lines = self._markdown_header(artifact, "画面笔记")
        lines.extend([artifact.visual_overview.strip() or "以下截图和解读用于补充理解视频中的画面。", ""])
        for note, image_path in images:
            label = self._visual_time_label(note)
            lines.extend([f"### {label}", ""])
            if image_path is not None:
                relative = image_path.relative_to(path.parent).as_posix()
                lines.extend([f"![{label} 的画面]({relative})", ""])
            else:
                lines.extend(["截图未能保存，以下仅保留已经取得的文字解读。", ""])
            lines.extend([note.description.strip() or "此画面尚未取得可靠解读。", ""])
        if not images:
            lines.append("未取得可用的画面截图。请查看处理提示，修复后可以重试。")
        self._write_text_atomic(path, "\n".join(lines).rstrip() + "\n")

    def _write_text_markdown(self, path: Path, artifact: ProcessingArtifact, heading: str, text: str) -> None:
        lines = self._markdown_header(artifact, heading)
        lines.append(text.strip())
        self._write_text_atomic(path, "\n".join(lines).rstrip() + "\n")

    def _markdown_header(self, artifact: ProcessingArtifact, heading: str) -> list[str]:
        lines = [
            f"# {artifact.title} - {heading}", "",
            f"- 来源：{artifact.source}",
            f"- 生成方式：{artifact.method}",
            f"- 识别语言：{artifact.language}",
            f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", "",
        ]
        if artifact.warnings:
            lines.extend(["## 处理提示", "", *[f"- {warning}" for warning in artifact.warnings], ""])
        lines.extend([f"## {heading}", ""])
        return lines

    def _write_text_atomic(self, path: Path, text: str) -> None:
        temporary: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                             prefix="." + path.name + ".", suffix=".tmp", delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(text)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    def _set_default_font(self, document: Document) -> None:
        style = document.styles["Normal"]
        style.font.name = "PingFang SC"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    def _paragraph(self, document: Document, text: str, style: str | None = None) -> None:
        run = document.add_paragraph(style=style).add_run(text)
        run.font.name = "PingFang SC"
        run.font.size = Pt(11)
        fonts = run._element.get_or_add_rPr().rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run._element.get_or_add_rPr().append(fonts)
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(key), "PingFang SC")
